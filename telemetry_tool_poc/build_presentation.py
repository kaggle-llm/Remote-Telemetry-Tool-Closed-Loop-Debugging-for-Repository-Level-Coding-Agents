import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
DESCRIPTOR = Path(r"C:\Users\hakan\AppData\Local\Temp\kilo\telemetry_presentation.json")
OUTPUT = ROOT / "telemetry_tool_poc" / "Telemetry_Lab_Prototype_TR.pptx"


def rgb(value):
    return RGBColor.from_string(value)


def add_text(slide, options):
    box = slide.shapes.add_textbox(
        Inches(options["x"]), Inches(options["y"]), Inches(options["w"]), Inches(options["h"])
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(options["text"].split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = options.get("fontFace", "Arial")
        paragraph.font.size = Pt(options.get("fontSize", 16))
        paragraph.font.bold = options.get("bold", False)
        paragraph.font.color.rgb = rgb(options.get("color", "202A27"))
        paragraph.alignment = {
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }.get(options.get("align", "left"), PP_ALIGN.LEFT)
        paragraph.space_after = Pt(4)
    return box


def add_shape(slide, options):
    shape_type = options.get("type", "RECT")
    if shape_type == "LINE":
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(options["x"]), Inches(options["y"]),
            Inches(options["w"]), Inches(options["h"]),
        )
    else:
        shape_id = MSO_SHAPE.ROUNDED_RECTANGLE if shape_type == "ROUNDRECT" else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(
            shape_id,
            Inches(options["x"]), Inches(options["y"]),
            Inches(options["w"]), Inches(options["h"]),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(options.get("fill", {}).get("color", "FFFFFF"))
    shape.line.color.rgb = rgb(options.get("line", {}).get("color", options.get("fill", {}).get("color", "FFFFFF")))
    return shape


def main():
    descriptor = json.loads(DESCRIPTOR.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for slide_data in descriptor["slides"]:
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(slide_data.get("background", {}).get("color", "FFFFFF"))
        for element in slide_data.get("elements", []):
            if element["type"] == "text":
                add_text(slide, element["options"])
            elif element["type"] == "shape":
                add_shape(slide, element["options"])
        if slide_data.get("notes"):
            slide.notes_slide.notes_text_frame.text = slide_data["notes"]

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
